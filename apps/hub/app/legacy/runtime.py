#!/usr/bin/env python3
from __future__ import annotations

import json
import hashlib
import base64
import io
import unicodedata
import ipaddress
import os
import re
import secrets
import socket
import sqlite3
import subprocess
import sys
import threading
import time
import traceback
import urllib.parse
import webbrowser
from datetime import datetime, timedelta, timezone
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

APP_NAME = "Order Recorder Hub"
APP_VERSION = "2.2.5"
DEFAULT_PORT = 17891
DEFAULT_RETENTION_DAYS = 7
DISCOVERY_PORT = 17892
DISCOVERY_MAGIC = b"ORDER_RECORDER_DISCOVER_V1"
VIETNAM_TZ = timezone(timedelta(hours=7))

# Dashboard/read access is intentionally limited to the Hub PC itself and
# devices connected through the same Tailscale tailnet. Ordinary LAN clients
# (for example the SUNMI) can still use the authenticated write APIs, but
# cannot read the dashboard/order database.
TAILSCALE_IPV4 = ipaddress.ip_network("100.64.0.0/10")
TAILSCALE_IPV6 = ipaddress.ip_network("fd7a:115c:a1e0::/48")


def dashboard_ip_allowed(value: str) -> bool:
    try:
        raw = str(value or "").split("%", 1)[0]
        addr = ipaddress.ip_address(raw)
        if addr.is_loopback:
            return True
        # Handle IPv4-mapped IPv6 addresses such as ::ffff:100.80.1.2.
        if isinstance(addr, ipaddress.IPv6Address) and addr.ipv4_mapped:
            addr = addr.ipv4_mapped
        if isinstance(addr, ipaddress.IPv4Address):
            return addr in TAILSCALE_IPV4
        return addr in TAILSCALE_IPV6
    except Exception:
        return False


def app_data_dir() -> Path:
    if os.name == "nt":
        base = os.environ.get("LOCALAPPDATA") or os.environ.get("APPDATA") or str(Path.home())
        return Path(base) / "OrderRecorderHub"
    return Path.home() / ".order-recorder-hub"


DATA_DIR = app_data_dir()
DATA_DIR.mkdir(parents=True, exist_ok=True)
CONFIG_PATH = DATA_DIR / "hub_config.json"
DB_PATH = DATA_DIR / "orders.db"
LOG_PATH = DATA_DIR / "hub.log"
BACKUP_DIR = DATA_DIR / "backups"
BACKUP_DIR.mkdir(parents=True, exist_ok=True)


def log(message: str) -> None:
    line = f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {message}"
    print(line, flush=True)
    try:
        with LOG_PATH.open("a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass



def _event_age_seconds(event_at: str) -> float:
    try:
        dt = to_vietnam_datetime(event_at)
        return max(0.0, (datetime.now(VIETNAM_TZ) - dt.astimezone(VIETNAM_TZ)).total_seconds())
    except Exception:
        return 10**9


def _pc_alert_payload(event: dict) -> tuple[str, str, str]:
    platform = str(event.get("platform") or "")
    label = "GrabFood" if platform == "grab" else "ShopeeFood" if platform == "shopeefood" else "Order Recorder"
    code = str(event.get("order_code") or "").strip()
    if not code:
        short_no = str(event.get("short_order_number") or "").strip()
        code = f"#{short_no}" if short_no else "đơn chưa xác định mã"
    typ = str(event.get("event_type") or "").lower()
    if typ == "risk":
        title = f"⚠ {label} có nguy cơ MISS SĐT"
        body = f"Đơn {code} chưa đọc được SĐT. Hãy kiểm tra thiết bị và backup ngay."
        severity = "risk"
    elif typ == "miss":
        title = f"⛔ {label} MISS SĐT"
        body = f"Đơn {code} đã hết giới hạn tự động. Hãy bổ sung SĐT thủ công nếu còn lấy được."
        severity = "miss"
    else:
        title = "Order Recorder Hub"
        body = str(event.get("reason") or "Cảnh báo hệ thống")
        severity = "risk"
    return title, body, severity


def _windows_alert_worker(title: str, body: str, severity: str) -> None:
    """Best-effort native Windows alert without third-party dependencies."""
    if os.name != "nt":
        log(f"PC alert (non-Windows test): {title} · {body}")
        return
    if PC_ALERT_SOUND:
        try:
            import winsound
            winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
        except Exception:
            pass
    try:
        enc=lambda x: base64.b64encode(str(x).encode("utf-8")).decode("ascii")
        title_b64,body_b64=enc(title),enc(body)
        url_b64=enc(f"http://127.0.0.1:{PORT}/#reconPanel")
        icon_kind = "Warning" if severity == "risk" else "Error"
        ps = f"""
$ErrorActionPreference='Stop'
$title=[Text.Encoding]::UTF8.GetString([Convert]::FromBase64String('{title_b64}'))
$body=[Text.Encoding]::UTF8.GetString([Convert]::FromBase64String('{body_b64}'))
$url=[Text.Encoding]::UTF8.GetString([Convert]::FromBase64String('{url_b64}'))
try {{
  [Windows.UI.Notifications.ToastNotificationManager, Windows.UI.Notifications, ContentType = WindowsRuntime] > $null
  [Windows.Data.Xml.Dom.XmlDocument, Windows.Data.Xml.Dom.XmlDocument, ContentType = WindowsRuntime] > $null
  $et=[Security.SecurityElement]::Escape($title)
  $eb=[Security.SecurityElement]::Escape($body)
  $eu=[Security.SecurityElement]::Escape($url)
  $xmlText='<toast activationType="protocol" launch="'+$eu+'"><visual><binding template="ToastGeneric"><text>'+ $et +'</text><text>'+ $eb +'</text></binding></visual><audio silent="true"/><actions><action content="Mở Hub" activationType="protocol" arguments="'+$eu+'"/></actions></toast>'
  $xml=New-Object Windows.Data.Xml.Dom.XmlDocument
  $xml.LoadXml($xmlText)
  $toast=[Windows.UI.Notifications.ToastNotification]::new($xml)
  $notifier=[Windows.UI.Notifications.ToastNotificationManager]::CreateToastNotifier('Order Recorder Hub')
  $notifier.Show($toast)
  exit 0
}} catch {{
  Add-Type -AssemblyName System.Windows.Forms
  Add-Type -AssemblyName System.Drawing
  $n=New-Object System.Windows.Forms.NotifyIcon
  $n.Icon=[System.Drawing.SystemIcons]::{icon_kind}
  $n.BalloonTipTitle=$title
  $n.BalloonTipText=$body
  $n.BalloonTipIcon=[System.Windows.Forms.ToolTipIcon]::{icon_kind}
  $n.Visible=$true
  $n.ShowBalloonTip(9000)
  Start-Sleep -Seconds 10
  $n.Dispose()
  exit 0
}}
"""
        encoded=base64.b64encode(ps.encode("utf-16le")).decode("ascii")
        flags=0x08000000
        r=subprocess.run(["powershell.exe","-NoProfile","-ExecutionPolicy","Bypass","-EncodedCommand",encoded],
                         stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,timeout=14,creationflags=flags)
        if r.returncode!=0:
            log(f"PC alert PowerShell trả mã {r.returncode}")
    except Exception as exc:
        log(f"PC alert không hiển thị được: {exc}")


def dispatch_pc_alert(event: dict, force: bool = False) -> bool:
    if not PC_ALERTS_ENABLED and not force:
        return False
    typ=str(event.get("event_type") or "").lower()
    if typ not in {"risk","miss"} and not force:
        return False
    if not force:
        age=_event_age_seconds(str(event.get("event_at") or ""))
        max_age=180.0 if typ=="risk" else 600.0
        if age>max_age:
            log(f"Bỏ native alert {typ}: event đã cũ {int(age)}s")
            return False
    title,body,severity=_pc_alert_payload(event)
    threading.Thread(target=_windows_alert_worker,args=(title,body,severity),daemon=True,name="hub-pc-alert").start()
    log(f"PC alert {typ or 'test'}: {title}")
    return True


def load_config() -> dict:
    cfg = {
        "port": DEFAULT_PORT,
        "retention_days": DEFAULT_RETENTION_DAYS,
        "api_key": secrets.token_urlsafe(24),
        "auto_open_dashboard": True,
        "hub_id": secrets.token_hex(8),
        "discovery_enabled": True,
        "admin_password_salt": "",
        "admin_password_hash": "",
        "backup_enabled": True,
        "backup_retention_days": DEFAULT_RETENTION_DAYS,
        "last_backup_date": "",
        "last_backup_at": "",
        "last_backup_file": "",
        "pc_alerts_enabled": True,
        "pc_alert_sound": True,
    }
    if CONFIG_PATH.exists():
        try:
            old = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
            if isinstance(old, dict):
                cfg.update({k: v for k, v in old.items() if k in cfg})
        except Exception as exc:
            log(f"Không đọc được config cũ: {exc}")
    CONFIG_PATH.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    return cfg


CONFIG = load_config()
PORT = int(CONFIG.get("port", DEFAULT_PORT))
RETENTION_DAYS = max(1, int(CONFIG.get("retention_days", DEFAULT_RETENTION_DAYS)))
API_KEY = str(CONFIG.get("api_key") or "")
HUB_ID = str(CONFIG.get("hub_id") or "")
DISCOVERY_ENABLED = bool(CONFIG.get("discovery_enabled", True))
BACKUP_ENABLED = bool(CONFIG.get("backup_enabled", True))
BACKUP_RETENTION_DAYS = min(RETENTION_DAYS, max(1, int(CONFIG.get("backup_retention_days", RETENTION_DAYS))))
PC_ALERTS_ENABLED = bool(CONFIG.get("pc_alerts_enabled", True))
PC_ALERT_SOUND = bool(CONFIG.get("pc_alert_sound", True))
# Backup chứa cùng dữ liệu cá nhân nên không được vượt retention nghiệp vụ.
CONFIG["backup_retention_days"] = BACKUP_RETENTION_DAYS


def db_connect() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=10, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    return conn


def init_db() -> None:
    with db_connect() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                platform TEXT NOT NULL,
                order_code TEXT NOT NULL DEFAULT '',
                short_order_number TEXT NOT NULL DEFAULT '',
                display_order_id TEXT NOT NULL DEFAULT '',
                full_order_id TEXT NOT NULL DEFAULT '',
                phone TEXT NOT NULL DEFAULT '',
                customer_name TEXT NOT NULL DEFAULT '',
                received_at TEXT NOT NULL,
                recorded_at TEXT NOT NULL,
                source_device TEXT NOT NULL DEFAULT '',
                sync_id TEXT NOT NULL DEFAULT '',
                agent_session_id TEXT NOT NULL DEFAULT '',
                dedup_key TEXT NOT NULL UNIQUE,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                manual_backup INTEGER NOT NULL DEFAULT 0,
                manual_backup_at TEXT NOT NULL DEFAULT '',
                deleted INTEGER NOT NULL DEFAULT 0,
                deleted_at TEXT NOT NULL DEFAULT '',
                deleted_by TEXT NOT NULL DEFAULT '',
                delete_reason TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_orders_received_at ON orders(received_at DESC);
            CREATE INDEX IF NOT EXISTS idx_orders_platform ON orders(platform);
            CREATE INDEX IF NOT EXISTS idx_orders_phone ON orders(phone);
            CREATE INDEX IF NOT EXISTS idx_orders_order_code ON orders(order_code);

            CREATE TABLE IF NOT EXISTS devices (
                source TEXT PRIMARY KEY,
                platform TEXT NOT NULL DEFAULT '',
                device_name TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL DEFAULT '',
                pending INTEGER NOT NULL DEFAULT 0,
                version TEXT NOT NULL DEFAULT '',
                last_seen TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                manual_backup INTEGER NOT NULL DEFAULT 0,
                manual_backup_at TEXT NOT NULL DEFAULT ''
            );
            CREATE INDEX IF NOT EXISTS idx_devices_platform ON devices(platform);

            CREATE TABLE IF NOT EXISTS order_audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                order_id INTEGER NOT NULL DEFAULT 0,
                order_code TEXT NOT NULL DEFAULT '',
                action TEXT NOT NULL,
                operator_name TEXT NOT NULL DEFAULT '',
                reason TEXT NOT NULL DEFAULT '',
                before_json TEXT NOT NULL DEFAULT '',
                after_json TEXT NOT NULL DEFAULT '',
                source_ip TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_audit_created_at ON order_audit_log(created_at DESC);
            CREATE INDEX IF NOT EXISTS idx_audit_order_id ON order_audit_log(order_id);

            CREATE TABLE IF NOT EXISTS order_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_key TEXT NOT NULL UNIQUE,
                platform TEXT NOT NULL,
                event_type TEXT NOT NULL,
                order_code TEXT NOT NULL DEFAULT '',
                short_order_number TEXT NOT NULL DEFAULT '',
                agent_session_id TEXT NOT NULL DEFAULT '',
                event_at TEXT NOT NULL,
                source_device TEXT NOT NULL DEFAULT '',
                reason TEXT NOT NULL DEFAULT '',
                attempt_count INTEGER NOT NULL DEFAULT 0,
                metadata_json TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_order_events_at ON order_events(event_at DESC);
            CREATE INDEX IF NOT EXISTS idx_order_events_platform ON order_events(platform);
            CREATE INDEX IF NOT EXISTS idx_order_events_type ON order_events(event_type);
            CREATE INDEX IF NOT EXISTS idx_order_events_session ON order_events(agent_session_id);
            CREATE INDEX IF NOT EXISTS idx_order_events_code ON order_events(order_code);
            """
        )
        # Migration an toàn cho database v2.0.0-v2.0.2 đang dùng thực tế.
        # SQLite CREATE TABLE IF NOT EXISTS không tự thêm cột mới vào bảng cũ.
        existing_cols = {str(r[1]) for r in conn.execute("PRAGMA table_info(orders)").fetchall()}
        if "manual_backup" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN manual_backup INTEGER NOT NULL DEFAULT 0")
        if "manual_backup_at" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN manual_backup_at TEXT NOT NULL DEFAULT ''")
        if "deleted" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN deleted INTEGER NOT NULL DEFAULT 0")
        if "deleted_at" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN deleted_at TEXT NOT NULL DEFAULT ''")
        if "deleted_by" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN deleted_by TEXT NOT NULL DEFAULT ''")
        if "delete_reason" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN delete_reason TEXT NOT NULL DEFAULT ''")
        if "agent_session_id" not in existing_cols:
            conn.execute("ALTER TABLE orders ADD COLUMN agent_session_id TEXT NOT NULL DEFAULT ''")
    cleanup_old_orders()


def parse_timestamp(value: str) -> datetime:
    s = str(value or "").strip()
    if not s:
        return datetime.now(VIETNAM_TZ)
    dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
    # Timestamp không có timezone được xem là giờ Việt Nam.
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=VIETNAM_TZ)
    return dt


def to_vietnam_datetime(value: str) -> datetime:
    return parse_timestamp(value).astimezone(VIETNAM_TZ)


def normalize_timestamp(value: str) -> str:
    return to_vietnam_datetime(value).isoformat(timespec="seconds")


def display_timestamp(value: str) -> str:
    try:
        return to_vietnam_datetime(value).strftime("%H:%M:%S %d-%m-%Y")
    except Exception:
        return str(value or "")


def local_now_iso() -> str:
    return datetime.now(VIETNAM_TZ).isoformat(timespec="seconds")


def date_part(value: str) -> str:
    if value:
        try:
            return to_vietnam_datetime(value).date().isoformat()
        except Exception:
            pass
    return datetime.now(VIETNAM_TZ).date().isoformat()


def normalize_phone(value: str) -> str:
    s = re.sub(r"[^0-9+]", "", str(value or ""))
    if s.startswith("+84"):
        s = "0" + s[3:]
    elif s.startswith("84") and len(s) >= 11:
        s = "0" + s[2:]
    digits = re.sub(r"\D", "", s)
    if len(digits) == 9 and digits[:1] in {"3", "5", "7", "8", "9"}:
        digits = "0" + digits
    return digits


def platform_name(value: str) -> str:
    v = str(value or "").strip().lower().replace(" ", "")
    if v in {"grab", "grabfood"}:
        return "grab"
    if v in {"shopeefood", "shopee", "spf"}:
        return "shopeefood"
    return ""


def clean_text(value, max_len=300) -> str:
    s = str(value or "").strip()
    return s[:max_len]


def normalize_grab_code(value: str) -> str:
    raw = clean_text(value, 120).upper().replace(" ", "")
    m = re.search(r"GF-?(\d+)", raw, re.I)
    return f"GF-{m.group(1)}" if m else raw


def extract_spf_display_code(display_id: str, order_code: str = "") -> str:
    for source in [display_id, order_code]:
        text = str(source or "")
        if text.upper().startswith("SPF-"):
            m = re.match(r"SPF-(\d+)$", text.strip(), re.I)
            if m:
                return m.group(1)
        matches = re.findall(r"#\s*(\d{3,10})", text)
        if matches:
            return matches[-1]
    return ""


def business_order_code(platform: str, order_code: str, display_id: str = "") -> str:
    if platform == "grab":
        return normalize_grab_code(order_code)
    if platform == "shopeefood":
        code = extract_spf_display_code(display_id, order_code)
        return f"SPF-{code}" if code else ""
    return clean_text(order_code, 120)


def build_dedup_key(order: dict) -> str:
    platform = order["platform"]
    full_id = order.get("full_order_id", "")
    if full_id:
        return f"{platform}:full:{full_id.lower()}"
    sync_id = order.get("sync_id", "")
    if sync_id:
        return f"{platform}:sync:{sync_id.lower()}"
    code = order.get("order_code") or order.get("display_order_id") or order.get("short_order_number")
    if code:
        return f"{platform}:day:{date_part(order.get('received_at',''))}:{code.lower()}"
    phone = order.get("phone", "")
    return f"{platform}:fallback:{date_part(order.get('received_at',''))}:{phone}:{order.get('received_at','')}"


def normalize_order(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    platform = platform_name(payload.get("platform"))
    if not platform:
        raise ValueError("platform phải là grab hoặc shopeefood")

    received_at = normalize_timestamp(clean_text(payload.get("receivedAt") or payload.get("received_at") or local_now_iso(), 80))
    recorded_at = normalize_timestamp(clean_text(payload.get("recordedAt") or payload.get("recorded_at") or local_now_iso(), 80))
    raw_order_code = clean_text(payload.get("orderCode") or payload.get("order_code"), 120)
    display_id = clean_text(payload.get("displayOrderId") or payload.get("display_order_id"), 160)
    short_no = clean_text(payload.get("shortOrderNumber") or payload.get("short_order_number"), 60)
    full_id = clean_text(payload.get("fullOrderId") or payload.get("full_order_id"), 220)
    phone = normalize_phone(payload.get("phone") or payload.get("receiverPhone") or payload.get("receiver_phone"))
    customer = clean_text(payload.get("customerName") or payload.get("customer_name"), 180)
    source = clean_text(payload.get("sourceDevice") or payload.get("source_device"), 120)
    sync_id = clean_text(payload.get("syncId") or payload.get("sync_id"), 180)
    agent_session_id = clean_text(payload.get("agentSessionId") or payload.get("agent_session_id") or payload.get("sessionId"), 220)

    order_code = business_order_code(platform, raw_order_code, display_id)
    if platform == "grab" and not order_code:
        order_code = normalize_grab_code(display_id or full_id)
    if platform == "shopeefood" and not order_code:
        order_code = ""
    if not any([order_code, display_id, short_no, full_id]):
        raise ValueError("Thiếu mã đơn")

    order = {
        "platform": platform,
        "order_code": order_code,
        "short_order_number": short_no,
        "display_order_id": display_id,
        "full_order_id": full_id,
        "phone": phone,
        "customer_name": customer,
        "received_at": received_at,
        "recorded_at": recorded_at,
        "source_device": source,
        "sync_id": sync_id,
        "agent_session_id": agent_session_id,
    }
    order["dedup_key"] = build_dedup_key(order)
    return order


def upsert_order(payload: dict) -> tuple[dict, bool]:
    order = normalize_order(payload)
    now = local_now_iso()
    with db_connect() as conn:
        existing = conn.execute("SELECT * FROM orders WHERE dedup_key=?", (order["dedup_key"],)).fetchone()
        # Nếu record cũ được tạo khi chưa có full ID, lần cập nhật sau phải gộp vào đúng đơn thay vì tạo trùng.
        if not existing and order.get("full_order_id"):
            day = date_part(order.get("received_at", ""))
            candidates = []
            if order.get("order_code"):
                candidates.append(("order_code", order["order_code"]))
            if order.get("display_order_id"):
                candidates.append(("display_order_id", order["display_order_id"]))
            if order.get("short_order_number"):
                candidates.append(("short_order_number", order["short_order_number"]))
            for field, value in candidates:
                row = conn.execute(
                    f"SELECT * FROM orders WHERE platform=? AND substr(received_at,1,10)=? AND {field}=? ORDER BY id DESC LIMIT 1",
                    (order["platform"], day, value),
                ).fetchone()
                if row:
                    existing = row
                    break
        if existing:
            # Không xóa dữ liệu tốt bằng payload rỗng. Payload mới có giá trị thì cập nhật.
            merged = dict(existing)
            for field in ["order_code", "short_order_number", "display_order_id", "full_order_id", "phone", "customer_name", "received_at", "recorded_at", "source_device", "sync_id", "agent_session_id"]:
                if order.get(field):
                    merged[field] = order[field]
            new_dedup_key = build_dedup_key({
                "platform": merged["platform"],
                "order_code": merged["order_code"],
                "display_order_id": merged["display_order_id"],
                "short_order_number": merged["short_order_number"],
                "full_order_id": merged["full_order_id"],
                "sync_id": merged["sync_id"],
                "phone": merged["phone"],
                "received_at": merged["received_at"],
            })
            conn.execute(
                """UPDATE orders SET order_code=?, short_order_number=?, display_order_id=?, full_order_id=?,
                   phone=?, customer_name=?, received_at=?, recorded_at=?, source_device=?, sync_id=?, agent_session_id=?, dedup_key=?, updated_at=?
                   WHERE id=?""",
                (
                    merged["order_code"], merged["short_order_number"], merged["display_order_id"], merged["full_order_id"],
                    merged["phone"], merged["customer_name"], merged["received_at"], merged["recorded_at"],
                    merged["source_device"], merged["sync_id"], merged.get("agent_session_id", ""), new_dedup_key, now, existing["id"]
                ),
            )
            row = conn.execute("SELECT * FROM orders WHERE id=?", (existing["id"],)).fetchone()
            return dict(row), False
        cur = conn.execute(
            """INSERT INTO orders(platform,order_code,short_order_number,display_order_id,full_order_id,phone,
               customer_name,received_at,recorded_at,source_device,sync_id,agent_session_id,dedup_key,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                order["platform"], order["order_code"], order["short_order_number"], order["display_order_id"],
                order["full_order_id"], order["phone"], order["customer_name"], order["received_at"], order["recorded_at"],
                order["source_device"], order["sync_id"], order.get("agent_session_id", ""), order["dedup_key"], now, now,
            ),
        )
        row = conn.execute("SELECT * FROM orders WHERE id=?", (cur.lastrowid,)).fetchone()
        return dict(row), True



class ManualBackupConflict(Exception):
    """Raised when staff tries to overwrite an existing different phone."""


def normalize_manual_order_code(platform: str, value: str) -> str:
    raw = clean_text(value, 120).strip().upper().replace(" ", "")
    if not raw:
        raise ValueError("Thiếu mã đơn")
    if platform == "grab":
        if raw.isdigit():
            raw = "GF-" + raw
        code = normalize_grab_code(raw)
        if not re.fullmatch(r"GF-\d{1,12}", code):
            raise ValueError("Mã Grab không hợp lệ. Ví dụ: GF-225")
        return code
    if platform == "shopeefood":
        m = re.fullmatch(r"(?:SPF-|#)?(\d{1,12})", raw, re.I)
        if not m:
            raise ValueError("Mã ShopeeFood không hợp lệ. Ví dụ: SPF-0686 hoặc #0686")
        return "SPF-" + m.group(1)
    raise ValueError("platform phải là grab hoặc shopeefood")


def validate_manual_phone(value: str) -> str:
    phone = normalize_phone(value)
    # Cho phép 10-11 chữ số bắt đầu bằng 0; vẫn hỗ trợ người dùng nhập +84/84
    # vì normalize_phone đã đổi về đầu 0 trước khi kiểm tra.
    if not re.fullmatch(r"0\d{9,10}", phone):
        raise ValueError("SĐT không hợp lệ. Hãy nhập SĐT Việt Nam 10-11 số")
    return phone


def validate_manual_received_at(value: str) -> str:
    received = normalize_timestamp(clean_text(value or local_now_iso(), 80))
    dt = to_vietnam_datetime(received)
    today = datetime.now(VIETNAM_TZ).date()
    oldest = today - timedelta(days=RETENTION_DAYS - 1)
    if dt.date() > today:
        raise ValueError("Thời gian nhận đơn không được ở tương lai")
    if dt.date() < oldest:
        raise ValueError(f"Hub chỉ lưu dữ liệu {RETENTION_DAYS} ngày gần nhất")
    return received


def _find_order_by_business_code(platform: str, code: str, received_at: str):
    day = date_part(received_at)
    with db_connect() as conn:
        rows = conn.execute(
            "SELECT * FROM orders WHERE deleted=0 AND platform=? AND substr(received_at,1,10)=? ORDER BY id DESC",
            (platform, day),
        ).fetchall()
        for row in rows:
            d = dict(row)
            current = business_order_code(
                d.get("platform", ""), d.get("order_code", ""), d.get("display_order_id", "")
            )
            if current.upper() == code.upper():
                return d
    return None


def manual_backup_order(payload: dict) -> tuple[dict, str]:
    """Insert a missed order or fill the phone of an existing incomplete order.

    Returns (row, action), where action is created, filled, or already_complete.
    It never overwrites a different existing phone automatically.
    """
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    platform = platform_name(payload.get("platform"))
    if not platform:
        raise ValueError("Hãy chọn Grab hoặc ShopeeFood")
    code = normalize_manual_order_code(platform, payload.get("orderCode") or payload.get("order_code"))
    phone = validate_manual_phone(payload.get("phone"))
    received_at = validate_manual_received_at(payload.get("receivedAt") or payload.get("received_at"))
    now = local_now_iso()
    agent_session_id = clean_text(payload.get("agentSessionId") or payload.get("agent_session_id"), 220)

    existing = _find_order_by_business_code(platform, code, received_at)
    if existing:
        old_phone = normalize_phone(existing.get("phone", ""))
        if old_phone:
            if old_phone == phone:
                return _row_for_api(existing), "already_complete"
            raise ManualBackupConflict(
                f"{code} đã có SĐT khác trong Hub. Không ghi đè tự động để tránh mất dữ liệu đúng."
            )
        with db_connect() as conn:
            conn.execute(
                """UPDATE orders SET phone=?, manual_backup=1, manual_backup_at=?,
                   agent_session_id=CASE WHEN ?<>'' THEN ? ELSE agent_session_id END, updated_at=? WHERE id=?""",
                (phone, now, agent_session_id, agent_session_id, now, int(existing["id"])),
            )
            row = conn.execute("SELECT * FROM orders WHERE id=?", (int(existing["id"]),)).fetchone()
        add_audit_log(int(existing["id"]), code, "FILL_BACKUP", "", "Bổ sung SĐT thủ công", existing, dict(row), "")
        log(f"MANUAL BACKUP: đã bổ sung SĐT cho {code} · platform={platform}")
        return _row_for_api(row), "filled"

    # Đơn miss hoàn toàn chưa có trong Hub: tạo record tối thiểu nhưng vẫn theo
    # cùng format/dedup logic để nếu agent sync được đơn đó về sau thì upsert có
    # thể gộp vào cùng record thay vì tạo bản sao.
    row, created = upsert_order({
        "platform": platform,
        "orderCode": code,
        "displayOrderId": code if platform == "shopeefood" else "",
        "phone": phone,
        "receivedAt": received_at,
        "recordedAt": now,
        "sourceDevice": "hub-manual-backup",
        "agentSessionId": agent_session_id,
    })
    with db_connect() as conn:
        conn.execute(
            "UPDATE orders SET manual_backup=1, manual_backup_at=?, updated_at=? WHERE id=?",
            (now, now, int(row["id"])),
        )
        saved = conn.execute("SELECT * FROM orders WHERE id=?", (int(row["id"]),)).fetchone()
    add_audit_log(int(saved["id"]), code, "CREATE_BACKUP", "", "Tạo đơn backup thủ công", None, dict(saved), "")
    log(f"MANUAL BACKUP: đã tạo {code} · platform={platform}")
    cleanup_old_orders()
    return _row_for_api(saved), "created" if created else "filled"


def save_config() -> None:
    CONFIG_PATH.write_text(json.dumps(CONFIG, ensure_ascii=False, indent=2), encoding="utf-8")


def _snapshot_for_audit(row) -> dict:
    if not row:
        return {}
    d = dict(row)
    keep = ["id", "platform", "order_code", "display_order_id", "full_order_id", "phone",
            "received_at", "recorded_at", "source_device", "manual_backup", "manual_backup_at",
            "deleted", "deleted_at", "deleted_by", "delete_reason", "agent_session_id"]
    return {k: d.get(k) for k in keep if k in d}


def add_audit_log(order_id: int, order_code: str, action: str, operator_name: str, reason: str,
                  before, after, source_ip: str) -> None:
    now = local_now_iso()
    before_json = json.dumps(_snapshot_for_audit(before), ensure_ascii=False, separators=(",", ":")) if before else ""
    after_json = json.dumps(_snapshot_for_audit(after), ensure_ascii=False, separators=(",", ":")) if after else ""
    with db_connect() as conn:
        conn.execute(
            """INSERT INTO order_audit_log(order_id,order_code,action,operator_name,reason,before_json,after_json,source_ip,created_at)
               VALUES(?,?,?,?,?,?,?,?,?)""",
            (int(order_id or 0), clean_text(order_code, 120), clean_text(action, 40), clean_text(operator_name, 120),
             clean_text(reason, 500), before_json, after_json, clean_text(source_ip, 80), now),
        )


def get_audit_logs(limit=100) -> list[dict]:
    try:
        limit = min(max(int(limit), 1), 500)
    except Exception:
        limit = 100
    with db_connect() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM order_audit_log ORDER BY id DESC LIMIT ?", (limit,)
        ).fetchall()]
    for r in rows:
        r["created_at_local"] = display_timestamp(r.get("created_at", ""))
        try:
            r["before"] = json.loads(r.get("before_json") or "{}")
        except Exception:
            r["before"] = {}
        try:
            r["after"] = json.loads(r.get("after_json") or "{}")
        except Exception:
            r["after"] = {}
        r.pop("before_json", None)
        r.pop("after_json", None)
    return rows


def _get_active_order(order_id: int):
    with db_connect() as conn:
        row = conn.execute("SELECT * FROM orders WHERE id=? AND deleted=0", (int(order_id),)).fetchone()
    return dict(row) if row else None


def _validate_operator_reason(payload: dict) -> tuple[str, str]:
    operator = clean_text(payload.get("operatorName") or payload.get("operator_name"), 120)
    reason = clean_text(payload.get("reason") or payload.get("note"), 500)
    if not operator:
        raise ValueError("Hãy nhập tên người thao tác")
    if not reason:
        raise ValueError("Hãy nhập lý do thao tác")
    return operator, reason


def edit_order(payload: dict, source_ip: str) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    try:
        order_id = int(payload.get("id"))
    except Exception:
        raise ValueError("ID đơn không hợp lệ")
    existing = _get_active_order(order_id)
    if not existing:
        raise ValueError("Không tìm thấy đơn hàng hoặc đơn đã bị xóa")
    operator, reason = _validate_operator_reason(payload)
    platform = platform_name(payload.get("platform") or existing.get("platform"))
    if not platform:
        raise ValueError("Hãy chọn GrabFood hoặc ShopeeFood")
    current_code = business_order_code(existing.get("platform", ""), existing.get("order_code", ""), existing.get("display_order_id", ""))
    code = normalize_manual_order_code(platform, payload.get("orderCode") or payload.get("order_code") or current_code)
    phone = validate_manual_phone(payload.get("phone"))
    received_at = validate_manual_received_at(payload.get("receivedAt") or payload.get("received_at") or existing.get("received_at"))
    now = local_now_iso()

    merged = dict(existing)
    merged["platform"] = platform
    merged["order_code"] = code
    if platform == "shopeefood":
        digits = code.split("-", 1)[1]
        merged["display_order_id"] = "#" + digits
    merged["phone"] = phone
    merged["received_at"] = received_at
    merged["updated_at"] = now
    new_dedup_key = build_dedup_key({
        "platform": merged["platform"],
        "order_code": merged.get("order_code", ""),
        "display_order_id": merged.get("display_order_id", ""),
        "short_order_number": merged.get("short_order_number", ""),
        "full_order_id": merged.get("full_order_id", ""),
        "sync_id": merged.get("sync_id", ""),
        "phone": merged.get("phone", ""),
        "received_at": merged.get("received_at", ""),
    })
    with db_connect() as conn:
        try:
            conn.execute(
                """UPDATE orders SET platform=?,order_code=?,display_order_id=?,phone=?,received_at=?,dedup_key=?,updated_at=? WHERE id=? AND deleted=0""",
                (merged["platform"], merged["order_code"], merged.get("display_order_id", ""), merged["phone"],
                 merged["received_at"], new_dedup_key, now, order_id),
            )
        except sqlite3.IntegrityError:
            raise ValueError("Thông tin sau khi sửa bị trùng với một đơn khác")
        row = conn.execute("SELECT * FROM orders WHERE id=?", (order_id,)).fetchone()
    add_audit_log(order_id, code, "EDIT", operator, reason, existing, dict(row), source_ip)
    log(f"ORDER EDIT: {code} · bởi {operator} · {reason}")
    return _row_for_api(row)


def _password_hash(password: str, salt_hex: str) -> str:
    salt = bytes.fromhex(salt_hex)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 240000)
    return digest.hex()


def admin_password_is_set() -> bool:
    return bool(str(CONFIG.get("admin_password_salt") or "") and str(CONFIG.get("admin_password_hash") or ""))


def verify_admin_password(password: str) -> bool:
    if not admin_password_is_set():
        return False
    try:
        candidate = _password_hash(str(password or ""), str(CONFIG.get("admin_password_salt") or ""))
        return secrets.compare_digest(candidate, str(CONFIG.get("admin_password_hash") or ""))
    except Exception:
        return False


def set_admin_password(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    current = str(payload.get("currentPassword") or "")
    new_password = str(payload.get("newPassword") or "")
    if len(new_password) < 8:
        raise ValueError("Mật khẩu admin phải có ít nhất 8 ký tự")
    if admin_password_is_set() and not verify_admin_password(current):
        raise PermissionError("Mật khẩu admin hiện tại không đúng")
    salt = secrets.token_hex(16)
    CONFIG["admin_password_salt"] = salt
    CONFIG["admin_password_hash"] = _password_hash(new_password, salt)
    save_config()
    log("ADMIN PASSWORD: đã đặt/đổi mật khẩu xóa đơn")
    return {"ok": True, "admin_password_set": True}


def soft_delete_order(payload: dict, source_ip: str) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    try:
        order_id = int(payload.get("id"))
    except Exception:
        raise ValueError("ID đơn không hợp lệ")
    operator, reason = _validate_operator_reason(payload)
    if not admin_password_is_set():
        raise PermissionError("Chưa đặt mật khẩu admin. Hãy đặt mật khẩu trong Cài đặt trên PC Hub trước")
    if not verify_admin_password(str(payload.get("adminPassword") or payload.get("admin_password") or "")):
        raise PermissionError("Mật khẩu admin không đúng")
    existing = _get_active_order(order_id)
    if not existing:
        raise ValueError("Không tìm thấy đơn hàng hoặc đơn đã bị xóa")
    code = business_order_code(existing.get("platform", ""), existing.get("order_code", ""), existing.get("display_order_id", "")) or existing.get("order_code", "")
    now = local_now_iso()
    tombstone_dedup = f"deleted:{order_id}:{existing.get('dedup_key','')}"
    with db_connect() as conn:
        conn.execute(
            """UPDATE orders SET deleted=1,deleted_at=?,deleted_by=?,delete_reason=?,dedup_key=?,updated_at=? WHERE id=? AND deleted=0""",
            (now, operator, reason, tombstone_dedup, now, order_id),
        )
        row = conn.execute("SELECT * FROM orders WHERE id=?", (order_id,)).fetchone()
    add_audit_log(order_id, code, "DELETE", operator, reason, existing, dict(row), source_ip)
    log(f"ORDER DELETE: {code} · bởi {operator} · {reason}")
    return {"id": order_id, "order_code": code, "deleted": True}


def _row_for_api(row) -> dict:
    d = dict(row)
    d["phone"] = normalize_phone(d.get("phone", ""))
    d["business_order_code"] = business_order_code(d.get("platform", ""), d.get("order_code", ""), d.get("display_order_id", ""))
    if d["business_order_code"]:
        d["order_code"] = d["business_order_code"]
    d["received_at_local"] = display_timestamp(d.get("received_at", ""))
    d["recorded_at_local"] = display_timestamp(d.get("recorded_at", ""))
    d["local_date"] = date_part(d.get("received_at", ""))
    d["manual_backup"] = bool(d.get("manual_backup", 0))
    d["deleted"] = bool(d.get("deleted", 0))
    d["platform_label"] = "GrabFood" if d.get("platform") == "grab" else "ShopeeFood" if d.get("platform") == "shopeefood" else d.get("platform", "")
    return d



def normalize_order_event(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON event phải là object")
    platform = platform_name(payload.get("platform"))
    if not platform:
        raise ValueError("platform event phải là grab hoặc shopeefood")
    raw_type = clean_text(payload.get("eventType") or payload.get("event_type") or payload.get("type"), 40).lower()
    aliases = {"order_seen": "seen", "seen": "seen", "detected": "seen",
               "order_captured": "captured", "captured": "captured", "completed": "captured",
               "order_risk": "risk", "risk": "risk", "at_risk": "risk", "warning": "risk",
               "order_miss": "miss", "miss": "miss", "needs_review": "miss"}
    event_type = aliases.get(raw_type, "")
    if not event_type:
        raise ValueError("eventType phải là seen, captured, risk hoặc miss")
    raw_code = clean_text(payload.get("orderCode") or payload.get("order_code"), 120)
    display_id = clean_text(payload.get("displayOrderId") or payload.get("display_order_id"), 160)
    order_code = business_order_code(platform, raw_code, display_id)
    if platform == "grab" and not order_code and raw_code:
        order_code = normalize_grab_code(raw_code)
    short_no = clean_text(payload.get("shortOrderNumber") or payload.get("short_order_number"), 80)
    session_id = clean_text(payload.get("agentSessionId") or payload.get("agent_session_id") or payload.get("sessionId"), 220)
    if not any([session_id, order_code, short_no]):
        raise ValueError("Event thiếu session/mã đơn")
    event_at = normalize_timestamp(clean_text(payload.get("eventAt") or payload.get("event_at") or payload.get("receivedAt") or local_now_iso(), 80))
    source = clean_text(payload.get("sourceDevice") or payload.get("source_device"), 140)
    reason = clean_text(payload.get("reason") or payload.get("failureReason") or payload.get("failure_reason"), 500)
    try:
        attempt_count = max(0, min(int(payload.get("attemptCount") or payload.get("attempt_count") or 0), 100))
    except Exception:
        attempt_count = 0
    metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
    metadata_json = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"))[:4000] if metadata else ""
    explicit = clean_text(payload.get("eventId") or payload.get("event_id"), 260)
    session_key = session_id or f"{date_part(event_at)}:{order_code or short_no}"
    event_key = explicit or f"{platform}:{session_key}:{event_type}"
    return {"event_key": event_key, "platform": platform, "event_type": event_type,
            "order_code": order_code, "short_order_number": short_no, "agent_session_id": session_id,
            "event_at": event_at, "source_device": source, "reason": reason,
            "attempt_count": attempt_count, "metadata_json": metadata_json}


def upsert_order_event(payload: dict) -> dict:
    e = normalize_order_event(payload)
    now = local_now_iso()
    created = False
    with db_connect() as conn:
        existed = conn.execute("SELECT id FROM order_events WHERE event_key=?", (e["event_key"],)).fetchone()
        created = existed is None
        conn.execute(
            """INSERT INTO order_events(event_key,platform,event_type,order_code,short_order_number,agent_session_id,
               event_at,source_device,reason,attempt_count,metadata_json,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?,?)
               ON CONFLICT(event_key) DO UPDATE SET
                 order_code=CASE WHEN excluded.order_code<>'' THEN excluded.order_code ELSE order_events.order_code END,
                 short_order_number=CASE WHEN excluded.short_order_number<>'' THEN excluded.short_order_number ELSE order_events.short_order_number END,
                 agent_session_id=CASE WHEN excluded.agent_session_id<>'' THEN excluded.agent_session_id ELSE order_events.agent_session_id END,
                 source_device=CASE WHEN excluded.source_device<>'' THEN excluded.source_device ELSE order_events.source_device END,
                 reason=CASE WHEN excluded.reason<>'' THEN excluded.reason ELSE order_events.reason END,
                 attempt_count=MAX(order_events.attempt_count, excluded.attempt_count),
                 metadata_json=CASE WHEN excluded.metadata_json<>'' THEN excluded.metadata_json ELSE order_events.metadata_json END""",
            (e["event_key"], e["platform"], e["event_type"], e["order_code"], e["short_order_number"],
             e["agent_session_id"], e["event_at"], e["source_device"], e["reason"], e["attempt_count"], e["metadata_json"], now),
        )
        row = conn.execute("SELECT * FROM order_events WHERE event_key=?", (e["event_key"],)).fetchone()
    out = _event_for_api(row)
    # Only a newly accepted RISK/MISS event may raise a PC notification. Retries/upserts never spam.
    if created and out.get("event_type") in {"risk", "miss"}:
        dispatch_pc_alert(out)
    return out


def _event_for_api(row) -> dict:
    d = dict(row)
    d["event_at_local"] = display_timestamp(d.get("event_at", ""))
    try:
        d["metadata"] = json.loads(d.get("metadata_json") or "{}")
    except Exception:
        d["metadata"] = {}
    d.pop("metadata_json", None)
    d["platform_label"] = "GrabFood" if d.get("platform") == "grab" else "ShopeeFood" if d.get("platform") == "shopeefood" else d.get("platform", "")
    return d


def _event_session_key(e: dict) -> str:
    session = str(e.get("agent_session_id") or "").strip()
    if session:
        return f"{e.get('platform','')}|{session}"
    code = str(e.get("order_code") or e.get("short_order_number") or "").strip().lower()
    return f"{e.get('platform','')}|{date_part(e.get('event_at',''))}|{code}"


def _orders_for_dates(start_date, end_date, include_deleted=False) -> list[dict]:
    sql = "SELECT * FROM orders WHERE 1=1" if include_deleted else "SELECT * FROM orders WHERE deleted=0"
    with db_connect() as conn:
        rows = [_row_for_api(r) for r in conn.execute(sql).fetchall()]
    out=[]
    for o in rows:
        try:
            d=to_vietnam_datetime(o.get("received_at", "")).date()
            if start_date <= d <= end_date:
                out.append(o)
        except Exception:
            pass
    return out


def reconciliation_for_range(start_value: str, end_value: str) -> dict:
    start_date, end_date = _parse_range_dates(start_value, end_value)
    with db_connect() as conn:
        raw = conn.execute("SELECT * FROM order_events ORDER BY event_at ASC,id ASC").fetchall()
    events=[]
    for r in raw:
        try:
            d=to_vietnam_datetime(r["event_at"]).date()
            if start_date <= d <= end_date:
                events.append(_event_for_api(r))
        except Exception:
            pass
    sessions={}
    for e in events:
        key=_event_session_key(e)
        x=sessions.setdefault(key,{"key":key,"platform":e.get("platform",""),"session_id":e.get("agent_session_id",""),
                                   "order_code":"","short_order_number":"","seen_at":"","captured_at":"","risk_at":"","miss_at":"",
                                   "risk_reason":"","miss_reason":"","attempt_count":0,"risk_attempt_count":0,"capture_mode":"","events":[]})
        x["events"].append(e)
        if e.get("order_code"): x["order_code"]=e["order_code"]
        if e.get("short_order_number"): x["short_order_number"]=e["short_order_number"]
        if e.get("agent_session_id"): x["session_id"]=e["agent_session_id"]
        typ=e.get("event_type")
        if typ=="seen" and not x["seen_at"]: x["seen_at"]=e.get("event_at","")
        elif typ=="captured":
            x["captured_at"]=e.get("event_at","")
            x["capture_mode"]=str((e.get("metadata") or {}).get("capture_mode") or "unknown").lower()
        elif typ=="risk":
            if not x["risk_at"]: x["risk_at"]=e.get("event_at","")
            x["risk_reason"]=e.get("reason","") or x["risk_reason"]
            x["risk_attempt_count"]=max(x["risk_attempt_count"],int(e.get("attempt_count") or 0))
        elif typ=="miss":
            x["miss_at"]=e.get("event_at","")
            x["miss_reason"]=e.get("reason","")
            x["attempt_count"]=max(x["attempt_count"],int(e.get("attempt_count") or 0))

    orders=_orders_for_dates(start_date,end_date,False)
    by_session={}
    by_code={}
    for o in orders:
        sid=str(o.get("agent_session_id") or "").strip()
        if sid: by_session[(o.get("platform"),sid)]=o
        code=str(o.get("business_order_code") or o.get("order_code") or "").upper()
        if code: by_code[(o.get("platform"),o.get("local_date"),code)]=o

    out_platforms={}
    issues=[]
    for p in ["grab","shopeefood"]:
        out_platforms[p]={"platform":p,"platform_label":"GrabFood" if p=="grab" else "ShopeeFood",
                          "seen":0,"captured":0,"agent_captured":0,"manual_rescue":0,"risk":0,"risk_active":0,"risk_recovered":0,
                          "miss":0,"pending":0,"data_complete":0,"manual_backup_recovered":0,"recovered_after_miss":0,"unresolved":0,
                          "capture_rate":None,"completion_rate":None,"telemetry_available":False}

    for x in sessions.values():
        # Ignore isolated captured/miss events only if a client never sent SEEN. They remain useful in timeline,
        # but the denominator must be based on explicit detection events.
        if not x.get("seen_at"):
            continue
        p=x["platform"]
        if p not in out_platforms: continue
        st=out_platforms[p]; st["telemetry_available"]=True; st["seen"]+=1
        if x.get("captured_at"):
            st["agent_captured"]+=1
            if x.get("capture_mode")=="manual": st["manual_rescue"]+=1
            else: st["captured"]+=1
        if x.get("risk_at"): st["risk"]+=1
        if x.get("miss_at"): st["miss"]+=1
        seen_date=date_part(x.get("seen_at"))
        order=None
        sid=str(x.get("session_id") or "")
        if sid: order=by_session.get((p,sid))
        if order is None and x.get("order_code"):
            order=by_code.get((p,seen_date,str(x["order_code"]).upper()))
        complete=bool(x.get("captured_at")) or bool(order and normalize_phone(order.get("phone","")))
        manual=bool(order and order.get("manual_backup") and normalize_phone(order.get("phone","")))
        if complete: st["data_complete"]+=1
        if manual and x.get("miss_at"): st["manual_backup_recovered"]+=1
        if complete and x.get("miss_at"): st["recovered_after_miss"]+=1
        if complete and x.get("risk_at"): st["risk_recovered"]+=1
        if not complete:
            st["unresolved"]+=1
            status="miss" if x.get("miss_at") else "risk" if x.get("risk_at") else "pending"
            if status=="risk": st["risk_active"]+=1
            elif status=="pending": st["pending"]+=1
            issues.append({"platform":p,"platform_label":st["platform_label"],"session_id":sid,
                           "order_id":int(order.get("id")) if order else 0,
                           "order_code":x.get("order_code") or ("#"+x.get("short_order_number") if x.get("short_order_number") else "Chưa xác định"),
                           "short_order_number":x.get("short_order_number",""),"seen_at":x.get("seen_at",""),
                           "seen_at_local":display_timestamp(x.get("seen_at","")),"status":status,
                           "risk_at":x.get("risk_at",""),"risk_at_local":display_timestamp(x.get("risk_at","")) if x.get("risk_at") else "",
                           "reason":x.get("miss_reason","") if status=="miss" else x.get("risk_reason","") if status=="risk" else "Đang chờ kết quả automation",
                           "attempt_count":x.get("attempt_count",0) or x.get("risk_attempt_count",0)})

    total={"seen":0,"captured":0,"agent_captured":0,"manual_rescue":0,"risk":0,"risk_active":0,"risk_recovered":0,
           "miss":0,"pending":0,"data_complete":0,"manual_backup_recovered":0,
           "recovered_after_miss":0,"unresolved":0,"capture_rate":None,"completion_rate":None,"telemetry_platforms":0}
    for st in out_platforms.values():
        if st["telemetry_available"]:
            total["telemetry_platforms"]+=1
            for k in ["seen","captured","agent_captured","manual_rescue","risk","risk_active","risk_recovered","miss","pending","data_complete","manual_backup_recovered","recovered_after_miss","unresolved"]:
                total[k]+=st[k]
            st["capture_rate"]=round(st["captured"]*100.0/st["seen"],2) if st["seen"] else None
            st["completion_rate"]=round(st["data_complete"]*100.0/st["seen"],2) if st["seen"] else None
    total["capture_rate"]=round(total["captured"]*100.0/total["seen"],2) if total["seen"] else None
    total["completion_rate"]=round(total["data_complete"]*100.0/total["seen"],2) if total["seen"] else None
    issues.sort(key=lambda x:x.get("seen_at", ""), reverse=True)
    return {"start_date":start_value,"end_date":end_value,"total":total,"platforms":out_platforms,"issues":issues[:500]}


def get_trash(limit=100, q="") -> list[dict]:
    try: limit=min(max(int(limit),1),500)
    except Exception: limit=100
    sql="SELECT * FROM orders WHERE deleted=1"
    params=[]
    if q:
        like="%"+clean_text(q,120)+"%"; digits=re.sub(r"\\D","",q)
        if digits:
            sql+=" AND (order_code LIKE ? OR display_order_id LIKE ? OR phone LIKE ? OR deleted_by LIKE ? OR delete_reason LIKE ?)"
            params.extend([like,like,"%"+digits+"%",like,like])
        else:
            sql+=" AND (order_code LIKE ? OR display_order_id LIKE ? OR deleted_by LIKE ? OR delete_reason LIKE ?)"
            params.extend([like,like,like,like])
    sql+=" ORDER BY deleted_at DESC,id DESC LIMIT ?"; params.append(limit)
    with db_connect() as conn:
        rows=[_row_for_api(r) for r in conn.execute(sql,params).fetchall()]
    for r in rows:
        r["deleted_at_local"]=display_timestamp(r.get("deleted_at","")) if r.get("deleted_at") else ""
    return rows


def restore_order(payload: dict, source_ip: str) -> dict:
    if not isinstance(payload,dict): raise ValueError("JSON phải là object")
    try: order_id=int(payload.get("id"))
    except Exception: raise ValueError("ID đơn không hợp lệ")
    operator,reason=_validate_operator_reason(payload)
    if not admin_password_is_set(): raise PermissionError("Chưa đặt mật khẩu admin")
    if not verify_admin_password(str(payload.get("adminPassword") or payload.get("admin_password") or "")):
        raise PermissionError("Mật khẩu admin không đúng")
    with db_connect() as conn:
        row=conn.execute("SELECT * FROM orders WHERE id=? AND deleted=1",(order_id,)).fetchone()
    if not row: raise ValueError("Không tìm thấy đơn trong Thùng rác")
    existing=dict(row)
    new_key=build_dedup_key(existing)
    with db_connect() as conn:
        collision=conn.execute("SELECT id FROM orders WHERE dedup_key=? AND deleted=0 AND id<>?",(new_key,order_id)).fetchone()
        if collision: raise ManualBackupConflict("Không thể khôi phục vì đã có một đơn đang hoạt động trùng dữ liệu")
        conn.execute("""UPDATE orders SET deleted=0,deleted_at='',deleted_by='',delete_reason='',dedup_key=?,updated_at=? WHERE id=?""",
                     (new_key,local_now_iso(),order_id))
        saved=conn.execute("SELECT * FROM orders WHERE id=?",(order_id,)).fetchone()
    code=business_order_code(existing.get("platform",""),existing.get("order_code",""),existing.get("display_order_id","")) or existing.get("order_code","")
    add_audit_log(order_id,code,"RESTORE",operator,reason,existing,dict(saved),source_ip)
    log(f"ORDER RESTORE: {code} · bởi {operator} · {reason}")
    return _row_for_api(saved)


def get_order_detail(order_id: int) -> dict:
    try: oid=int(order_id)
    except Exception: raise ValueError("ID đơn không hợp lệ")
    with db_connect() as conn:
        row=conn.execute("SELECT * FROM orders WHERE id=?",(oid,)).fetchone()
        audits=[dict(r) for r in conn.execute("SELECT * FROM order_audit_log WHERE order_id=? ORDER BY created_at ASC,id ASC",(oid,)).fetchall()]
    if not row: raise ValueError("Không tìm thấy đơn hàng")
    order=_row_for_api(row)
    sid=str(order.get("agent_session_id") or "").strip()
    code=str(order.get("business_order_code") or order.get("order_code") or "").strip()
    with db_connect() as conn:
        if sid:
            evrows=conn.execute("SELECT * FROM order_events WHERE agent_session_id=? ORDER BY event_at ASC,id ASC",(sid,)).fetchall()
        elif code:
            # Business display codes can be reused on another day; legacy fallback must stay date-scoped.
            evrows=conn.execute("SELECT * FROM order_events WHERE platform=? AND order_code=? AND substr(event_at,1,10)=? ORDER BY event_at ASC,id ASC",
                                (order.get("platform"),code,order.get("local_date",""))).fetchall()
        else:
            evrows=[]
    events=[_event_for_api(r) for r in evrows]
    timeline=[]
    timeline.append({"at":order.get("created_at",""),"at_local":display_timestamp(order.get("created_at","")),"kind":"system","type":"created","label":"Record được tạo trong Hub","detail":order.get("source_device","")})
    labels={"seen":"Agent phát hiện đơn","captured":"Automation lấy được SĐT","risk":"Cảnh báo sớm: nguy cơ MISS","miss":"Automation báo MISS"}
    for e in events:
        detail=e.get("reason","")
        if e.get("event_type")=="captured":
            mode=str((e.get("metadata") or {}).get("capture_mode") or "unknown")
            detail=(detail+(" · " if detail else "")+"mode="+mode)
        if e.get("attempt_count"): detail=(detail+(" · " if detail else "")+f"attempt={e['attempt_count']}")
        timeline.append({"at":e.get("event_at",""),"at_local":e.get("event_at_local",""),"kind":"event","type":e.get("event_type",""),"label":labels.get(e.get("event_type"),e.get("event_type","")),"detail":detail,"source":e.get("source_device","")})
    action_labels={"CREATE_BACKUP":"Tạo đơn thủ công","FILL_BACKUP":"Bổ sung SĐT thủ công","EDIT":"Chỉnh sửa đơn","DELETE":"Xóa vào Thùng rác","RESTORE":"Khôi phục từ Thùng rác","IMPORT_CREATE":"Import tạo đơn mới","IMPORT_FILL":"Import bổ sung SĐT"}
    audit_out=[]
    for a in audits:
        try:a["before"]=json.loads(a.get("before_json") or "{}")
        except Exception:a["before"]={}
        try:a["after"]=json.loads(a.get("after_json") or "{}")
        except Exception:a["after"]={}
        a["created_at_local"]=display_timestamp(a.get("created_at",""))
        a.pop("before_json",None);a.pop("after_json",None);audit_out.append(a)
        detail=""
        if a.get("operator_name"): detail=f"Bởi {a['operator_name']}"
        if a.get("reason"): detail+=(" · " if detail else "")+a["reason"]
        timeline.append({"at":a.get("created_at",""),"at_local":a.get("created_at_local",""),"kind":"audit","type":a.get("action",""),"label":action_labels.get(a.get("action"),a.get("action","")),"detail":detail})
    timeline.sort(key=lambda x:x.get("at", ""))
    return {"order":order,"events":events,"audit_logs":audit_out,"timeline":timeline}


def _cleanup_database_backups() -> int:
    cutoff=time.time()-BACKUP_RETENTION_DAYS*86400
    removed=0
    try:
        for f in BACKUP_DIR.glob("orders_*.db"):
            try:
                if f.stat().st_mtime < cutoff:
                    f.unlink();removed+=1
            except Exception:
                pass
    except Exception:
        pass
    return removed


def create_database_backup(force=False) -> dict:
    if not BACKUP_ENABLED and not force:
        return {"ok":False,"skipped":True,"reason":"backup_disabled"}
    today=datetime.now(VIETNAM_TZ).date().isoformat()
    if not force and str(CONFIG.get("last_backup_date") or "")==today:
        return {"ok":True,"skipped":True,"file":str(CONFIG.get("last_backup_file") or ""),"at":str(CONFIG.get("last_backup_at") or "")}
    stamp=datetime.now(VIETNAM_TZ).strftime("%Y-%m-%d_%H%M%S")
    final=BACKUP_DIR/f"orders_{stamp}.db"
    tmp=BACKUP_DIR/f".orders_{stamp}.tmp"
    src=dst=None
    try:
        src=sqlite3.connect(DB_PATH,timeout=15)
        dst=sqlite3.connect(tmp,timeout=15)
        src.backup(dst)
        dst.commit();dst.close();dst=None;src.close();src=None
        os.replace(tmp,final)
        now=local_now_iso()
        CONFIG["last_backup_date"]=today;CONFIG["last_backup_at"]=now;CONFIG["last_backup_file"]=str(final);save_config()
        removed=_cleanup_database_backups()
        log(f"DATABASE BACKUP: {final.name}"+(f" · dọn {removed} bản cũ" if removed else ""))
        return {"ok":True,"skipped":False,"file":str(final),"name":final.name,"at":now,"retention_days":BACKUP_RETENTION_DAYS}
    except Exception as exc:
        log(f"DATABASE BACKUP lỗi: {exc}")
        try:
            if tmp.exists(): tmp.unlink()
        except Exception: pass
        return {"ok":False,"skipped":False,"error":str(exc)}
    finally:
        try:
            if dst is not None: dst.close()
        except Exception: pass
        try:
            if src is not None: src.close()
        except Exception: pass


def database_backup_status() -> dict:
    files=[]
    try:
        for f in sorted(BACKUP_DIR.glob("orders_*.db"),key=lambda x:x.stat().st_mtime,reverse=True)[:10]:
            files.append({"name":f.name,"path":str(f),"size_bytes":f.stat().st_size,"modified_at":datetime.fromtimestamp(f.stat().st_mtime,VIETNAM_TZ).isoformat(timespec="seconds")})
    except Exception: pass
    return {"enabled":BACKUP_ENABLED,"backup_dir":str(BACKUP_DIR),"retention_days":BACKUP_RETENTION_DAYS,
            "last_backup_at":str(CONFIG.get("last_backup_at") or ""),"last_backup_file":str(CONFIG.get("last_backup_file") or ""),"recent":files}



# ---------------------------------------------------------------------------
# Password-protected XLSX import
# ---------------------------------------------------------------------------
IMPORT_XLSX_MAX_BYTES = 5 * 1024 * 1024
IMPORT_XLSX_MAX_ROWS = 10000


def _import_header_key(value) -> str:
    text = str(value or "").strip().lower()
    text = text.replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    aliases = {
        "ma don hang": "order_code", "ma don": "order_code", "order code": "order_code", "ordercode": "order_code",
        "sdt": "phone", "so dien thoai": "phone", "dien thoai": "phone", "phone": "phone", "phone number": "phone",
        "thoi gian nhan don": "received_at", "thoi gian": "received_at", "received at": "received_at", "receivedat": "received_at",
        "nen tang": "platform", "platform": "platform", "app": "platform",
    }
    return aliases.get(text, "")


def _import_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_import_received_at(value) -> str:
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=VIETNAM_TZ)
        return validate_manual_received_at(dt.astimezone(VIETNAM_TZ).isoformat(timespec="seconds"))
    # openpyxl can return date objects for date-only cells.
    try:
        from datetime import date as _date
        if isinstance(value, _date):
            dt = datetime(value.year, value.month, value.day, 0, 0, 0, tzinfo=VIETNAM_TZ)
            return validate_manual_received_at(dt.isoformat(timespec="seconds"))
    except Exception:
        pass
    raw = _import_cell_text(value)
    if not raw:
        raise ValueError("Thiếu Thời gian nhận đơn")
    # First accept ISO formats handled by the Hub itself.
    try:
        return validate_manual_received_at(raw)
    except Exception:
        pass
    formats = [
        "%H:%M:%S %d-%m-%Y", "%H:%M %d-%m-%Y",
        "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
        "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
    ]
    for fmt in formats:
        try:
            dt = datetime.strptime(raw, fmt).replace(tzinfo=VIETNAM_TZ)
            return validate_manual_received_at(dt.isoformat(timespec="seconds"))
        except Exception:
            continue
    raise ValueError("Thời gian nhận đơn không hợp lệ")


def _infer_import_platform(raw_platform, raw_code) -> str:
    p = platform_name(raw_platform)
    if p:
        return p
    code = _import_cell_text(raw_code).upper().replace(" ", "")
    if code.startswith("GF-"):
        return "grab"
    if code.startswith("SPF-") or code.startswith("#"):
        return "shopeefood"
    raise ValueError("Không xác định được nền tảng. Hãy dùng mã GF-/SPF- hoặc thêm cột Nền tảng")


def _find_any_order_by_business_code(platform: str, code: str, received_at: str):
    day = date_part(received_at)
    with db_connect() as conn:
        rows = conn.execute(
            "SELECT * FROM orders WHERE platform=? AND substr(received_at,1,10)=? ORDER BY deleted ASC,id DESC",
            (platform, day),
        ).fetchall()
        for row in rows:
            d = dict(row)
            current = business_order_code(d.get("platform", ""), d.get("order_code", ""), d.get("display_order_id", ""))
            if current.upper() == code.upper():
                return d
    return None


def import_xlsx_orders(payload: dict, source_ip: str) -> dict:
    """Import Hub order rows from a .xlsx file.

    Security / safety rules:
    - requires existing admin password and a correct password on every import;
    - requires operator name + reason;
    - creates a SQLite online backup before changing any row;
    - never overwrites an existing different phone automatically;
    - never silently resurrects a row in Trash.
    """
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    operator, reason = _validate_operator_reason(payload)
    if not admin_password_is_set():
        raise PermissionError("Chưa đặt mật khẩu admin. Hãy đặt mật khẩu trong Cài đặt trên PC Hub trước")
    if not verify_admin_password(str(payload.get("adminPassword") or payload.get("admin_password") or "")):
        raise PermissionError("Mật khẩu admin không đúng")

    filename = clean_text(payload.get("filename"), 220)
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Chỉ hỗ trợ file Excel .xlsx")
    encoded = str(payload.get("contentBase64") or payload.get("content_base64") or "")
    if not encoded:
        raise ValueError("Chưa chọn file .xlsx")
    try:
        raw = base64.b64decode(encoded, validate=True)
    except Exception:
        raise ValueError("File upload không hợp lệ")
    if not raw or len(raw) > IMPORT_XLSX_MAX_BYTES:
        raise ValueError("File .xlsx tối đa 5 MB")

    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Không đọc được file .xlsx: {exc}")

    sheet = None
    header_row = None
    columns = None
    # Search the first 10 rows of each sheet so files with a title above the
    # table still work, while requiring a recognizable order-code header.
    for ws in wb.worksheets:
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
            found = {}
            for col_idx, value in enumerate(row):
                key = _import_header_key(value)
                if key and key not in found:
                    found[key] = col_idx
            if "order_code" in found:
                sheet, header_row, columns = ws, idx, found
                break
        if sheet is not None:
            break
    if sheet is None:
        raise ValueError("Không tìm thấy cột Mã đơn hàng trong file")
    if "received_at" not in columns:
        raise ValueError("File phải có cột Thời gian nhận đơn")

    safety_backup = create_database_backup(force=True)
    if not safety_backup.get("ok"):
        raise RuntimeError("Không thể tạo backup an toàn trước khi import: " + str(safety_backup.get("error") or "unknown"))

    result = {
        "ok": True, "filename": filename, "sheet": sheet.title, "header_row": header_row,
        "total_rows": 0, "created": 0, "updated": 0, "skipped": 0,
        "conflicts": 0, "invalid": 0, "conflict_rows": [], "invalid_rows": [],
        "safety_backup": {"name": safety_backup.get("name", ""), "at": safety_backup.get("at", "")},
    }
    now = local_now_iso()
    processed = 0
    for excel_row, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        if processed >= IMPORT_XLSX_MAX_ROWS:
            result["invalid"] += 1
            result["invalid_rows"].append({"row": excel_row, "reason": f"Vượt giới hạn {IMPORT_XLSX_MAX_ROWS} dòng/import"})
            break
        values = list(row)
        def getcol(name):
            pos = columns.get(name)
            return values[pos] if pos is not None and pos < len(values) else None
        raw_code = getcol("order_code")
        raw_phone = getcol("phone")
        raw_time = getcol("received_at")
        raw_platform = getcol("platform")
        if not any(_import_cell_text(x) for x in (raw_code, raw_phone, raw_time, raw_platform)):
            continue
        processed += 1
        result["total_rows"] += 1
        try:
            platform = _infer_import_platform(raw_platform, raw_code)
            code = normalize_manual_order_code(platform, _import_cell_text(raw_code))
            received_at = _parse_import_received_at(raw_time)
            phone_text = _import_cell_text(raw_phone)
            phone = validate_manual_phone(phone_text) if phone_text else ""

            existing = _find_any_order_by_business_code(platform, code, received_at)
            if existing:
                if int(existing.get("deleted") or 0):
                    result["conflicts"] += 1
                    if len(result["conflict_rows"]) < 100:
                        result["conflict_rows"].append({"row": excel_row, "code": code, "reason": "Đơn đang ở Thùng rác; hãy khôi phục thay vì import lại"})
                    continue
                old_phone = normalize_phone(existing.get("phone", ""))
                if not phone or old_phone == phone:
                    result["skipped"] += 1
                    continue
                if old_phone and old_phone != phone:
                    result["conflicts"] += 1
                    if len(result["conflict_rows"]) < 100:
                        result["conflict_rows"].append({"row": excel_row, "code": code, "reason": "Hub đã có SĐT khác; không ghi đè tự động"})
                    continue
                before = dict(existing)
                with db_connect() as conn:
                    conn.execute("UPDATE orders SET phone=?,updated_at=? WHERE id=? AND deleted=0", (phone, now, int(existing["id"])))
                    saved = conn.execute("SELECT * FROM orders WHERE id=?", (int(existing["id"]),)).fetchone()
                audit_reason = f"{reason} · Import {filename} · dòng {excel_row}"
                add_audit_log(int(existing["id"]), code, "IMPORT_FILL", operator, audit_reason, before, dict(saved), source_ip)
                result["updated"] += 1
                continue

            display_id = "#" + code.split("-", 1)[1] if platform == "shopeefood" else ""
            saved, created = upsert_order({
                "platform": platform,
                "orderCode": code,
                "displayOrderId": display_id,
                "phone": phone,
                "receivedAt": received_at,
                "recordedAt": now,
                "sourceDevice": "hub-xlsx-import",
            })
            if not created:
                # Conservative fallback: if upsert matched something by an
                # alternate key, do not report a new record incorrectly.
                result["skipped"] += 1
                continue
            audit_reason = f"{reason} · Import {filename} · dòng {excel_row}"
            add_audit_log(int(saved["id"]), code, "IMPORT_CREATE", operator, audit_reason, None, dict(saved), source_ip)
            result["created"] += 1
        except Exception as exc:
            result["invalid"] += 1
            if len(result["invalid_rows"]) < 100:
                result["invalid_rows"].append({"row": excel_row, "code": _import_cell_text(raw_code), "reason": str(exc)})

    try:
        wb.close()
    except Exception:
        pass
    cleanup_old_orders()
    log(f"XLSX IMPORT: {filename} · bởi {operator} · total={result['total_rows']} created={result['created']} updated={result['updated']} skipped={result['skipped']} conflicts={result['conflicts']} invalid={result['invalid']}")
    return result

def cleanup_old_orders() -> int:
    cutoff = datetime.now(VIETNAM_TZ) - timedelta(days=RETENTION_DAYS)
    try:
        with db_connect() as conn:
            rows = conn.execute("SELECT id, received_at FROM orders").fetchall()
            ids = []
            for r in rows:
                try:
                    if to_vietnam_datetime(r["received_at"]) < cutoff:
                        ids.append(int(r["id"]))
                except Exception:
                    pass
            if ids:
                conn.executemany("DELETE FROM orders WHERE id=?", [(i,) for i in ids])
            event_rows = conn.execute("SELECT id,event_at FROM order_events").fetchall()
            event_ids=[]
            for r in event_rows:
                try:
                    if to_vietnam_datetime(r["event_at"]) < cutoff:
                        event_ids.append(int(r["id"]))
                except Exception:
                    pass
            if event_ids:
                conn.executemany("DELETE FROM order_events WHERE id=?", [(i,) for i in event_ids])
            return len(ids)
    except Exception as exc:
        log(f"Cleanup lỗi: {exc}")
        return 0


def _query_candidates(platform="", q="") -> list[dict]:
    sql = "SELECT * FROM orders WHERE deleted=0"
    params = []
    if platform in {"grab", "shopeefood"}:
        sql += " AND platform=?"
        params.append(platform)
    if q:
        like = "%" + q.strip() + "%"
        raw_digits = re.sub(r"\D", "", q)
        if raw_digits:
            digits = "%" + raw_digits + "%"
            sql += " AND (order_code LIKE ? OR display_order_id LIKE ? OR full_order_id LIKE ? OR short_order_number LIKE ? OR phone LIKE ?)"
            params.extend([like, like, like, like, digits])
        else:
            sql += " AND (order_code LIKE ? OR display_order_id LIKE ? OR full_order_id LIKE ? OR short_order_number LIKE ?)"
            params.extend([like, like, like, like])
    sql += " ORDER BY id DESC LIMIT 10000"
    with db_connect() as conn:
        rows = [_row_for_api(r) for r in conn.execute(sql, params).fetchall()]
    rows.sort(key=lambda d: to_vietnam_datetime(d.get("received_at", "")), reverse=True)
    return rows


def query_orders(days=1, platform="", q="") -> list[dict]:
    days = min(max(int(days), 1), RETENTION_DAYS)
    today = datetime.now(VIETNAM_TZ).date()
    start_date = today - timedelta(days=days - 1)
    return [o for o in _query_candidates(platform, q)
            if start_date <= to_vietnam_datetime(o.get("received_at", "")).date() <= today][:5000]


def query_orders_date(date_value: str, platform="", q="") -> list[dict]:
    date_value = clean_text(date_value, 10)
    try:
        selected = datetime.strptime(date_value, "%Y-%m-%d").date()
    except Exception:
        raise ValueError("Ngày không hợp lệ, định dạng phải là YYYY-MM-DD")

    today = datetime.now(VIETNAM_TZ).date()
    oldest = today - timedelta(days=RETENTION_DAYS - 1)
    if selected > today:
        raise ValueError("Không thể chọn ngày trong tương lai")
    if selected < oldest:
        raise ValueError(f"Hub chỉ lưu dữ liệu {RETENTION_DAYS} ngày gần nhất")

    return [o for o in _query_candidates(platform, q)
            if to_vietnam_datetime(o.get("received_at", "")).date() == selected][:5000]


def _parse_range_dates(start_value: str, end_value: str):
    start_value = clean_text(start_value, 10)
    end_value = clean_text(end_value, 10)
    try:
        start_date = datetime.strptime(start_value, "%Y-%m-%d").date()
        end_date = datetime.strptime(end_value, "%Y-%m-%d").date()
    except Exception:
        raise ValueError("Ngày không hợp lệ, định dạng phải là YYYY-MM-DD")

    today = datetime.now(VIETNAM_TZ).date()
    oldest = today - timedelta(days=RETENTION_DAYS - 1)
    if start_date > end_date:
        raise ValueError("Từ ngày không được lớn hơn Đến ngày")
    if end_date > today:
        raise ValueError("Không thể chọn ngày trong tương lai")
    if start_date < oldest:
        raise ValueError(f"Hub chỉ lưu dữ liệu {RETENTION_DAYS} ngày gần nhất")
    return start_date, end_date


def query_orders_range(start_value: str, end_value: str, platform="", q="") -> list[dict]:
    start_date, end_date = _parse_range_dates(start_value, end_value)
    return [o for o in _query_candidates(platform, q)
            if start_date <= to_vietnam_datetime(o.get("received_at", "")).date() <= end_date][:5000]


def stats_today() -> dict:
    today = datetime.now(VIETNAM_TZ).date()
    out = {"grab": 0, "shopeefood": 0, "total": 0, "complete": 0}
    with db_connect() as conn:
        rows = conn.execute("SELECT platform, phone, received_at FROM orders WHERE deleted=0").fetchall()
    for r in rows:
        try:
            if to_vietnam_datetime(r["received_at"]).date() != today:
                continue
        except Exception:
            continue
        platform = r["platform"]
        if platform in out:
            out[platform] += 1
        out["total"] += 1
        if str(r["phone"] or "").strip():
            out["complete"] += 1
    return out


DEVICE_ONLINE_SECONDS = 95
HUB_STARTED_AT = local_now_iso()



def stats_for_date(date_value: str) -> dict:
    orders = query_orders_date(date_value)
    out = {"grab": 0, "shopeefood": 0, "total": 0, "complete": 0}
    for o in orders:
        p = o.get("platform", "")
        if p in out:
            out[p] += 1
        out["total"] += 1
        if normalize_phone(o.get("phone", "")):
            out["complete"] += 1
    return out

def stats_for_range(start_value: str, end_value: str) -> dict:
    orders = query_orders_range(start_value, end_value)
    out = {"grab": 0, "shopeefood": 0, "total": 0, "complete": 0}
    for o in orders:
        p = o.get("platform", "")
        if p in out:
            out[p] += 1
        out["total"] += 1
        if normalize_phone(o.get("phone", "")):
            out["complete"] += 1
    return out


def upsert_heartbeat(payload: dict) -> dict:
    if not isinstance(payload, dict):
        raise ValueError("JSON phải là object")
    source = clean_text(payload.get("source"), 80).lower()
    platform = platform_name(payload.get("platform"))
    device_name = clean_text(payload.get("deviceName") or payload.get("device_name"), 140)
    status = clean_text(payload.get("status") or "ready", 80)
    version = clean_text(payload.get("version"), 40)
    try:
        pending = max(0, min(int(payload.get("pending") or 0), 100000))
    except Exception:
        pending = 0
    if not source:
        raise ValueError("Thiếu source heartbeat")
    if platform not in {"grab", "shopeefood"}:
        raise ValueError("platform heartbeat phải là grab hoặc shopeefood")
    now = local_now_iso()
    with db_connect() as conn:
        conn.execute(
            """INSERT INTO devices(source,platform,device_name,status,pending,version,last_seen,updated_at)
               VALUES(?,?,?,?,?,?,?,?)
               ON CONFLICT(source) DO UPDATE SET
                 platform=excluded.platform, device_name=excluded.device_name, status=excluded.status,
                 pending=excluded.pending, version=excluded.version, last_seen=excluded.last_seen,
                 updated_at=excluded.updated_at""",
            (source, platform, device_name, status, pending, version, now, now),
        )
    return {"source": source, "platform": platform, "device_name": device_name, "status": status,
            "pending": pending, "version": version, "last_seen": now}


def device_status() -> dict:
    now = datetime.now().astimezone()
    with db_connect() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM devices ORDER BY last_seen DESC").fetchall()]

    def latest(platform: str) -> dict:
        for r in rows:
            if r.get("platform") != platform:
                continue
            try:
                seen = datetime.fromisoformat(str(r.get("last_seen", "")).replace("Z", "+00:00")).astimezone()
                age = max(0, int((now - seen).total_seconds()))
            except Exception:
                age = 999999
            online = age <= DEVICE_ONLINE_SECONDS
            status = r.get("status", "")
            if not online:
                message = "Mất kết nối"
            elif status == "ready":
                message = "Hoạt động"
            elif status == "auto_off":
                message = "Auto đang tắt"
            elif status == "recording_off":
                message = "Ghi nhận đang tắt"
            else:
                message = status or "Đã kết nối"
            return {
                "online": online,
                "status": status,
                "pending": int(r.get("pending") or 0),
                "device_name": r.get("device_name", ""),
                "version": r.get("version", ""),
                "last_seen": r.get("last_seen", ""),
                "age_seconds": age,
                "message": message,
            }
        return {"online": False, "status": "never", "pending": 0, "device_name": "", "version": "",
                "last_seen": "", "age_seconds": None, "message": "Chưa kết nối"}

    return {
        "hub": {"online": True, "status": "ready", "message": "Hoạt động", "version": APP_VERSION, "started_at": HUB_STARTED_AT},
        "grab": latest("grab"),
        "shopeefood": latest("shopeefood"),
    }


def get_local_ip() -> str:
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        try:
            return socket.gethostbyname(socket.gethostname())
        except Exception:
            return "127.0.0.1"


def discovery_payload() -> dict:
    """Public, non-secret identity used only to locate this Hub on the local network."""
    ip = get_local_ip()
    return {
        "service": "order-recorder-hub",
        "protocol": 1,
        "hub_id": HUB_ID,
        "app": APP_NAME,
        "version": APP_VERSION,
        "lan_ip": ip,
        "port": PORT,
        "url": f"http://{ip}:{PORT}",
    }


def discovery_server():
    """
    UDP discovery responder for SUNMI.

    Client broadcasts ORDER_RECORDER_DISCOVER_V1 to UDP/17892. Hub replies only
    to the sender with non-secret location metadata. The client must still verify
    the candidate by sending an authenticated heartbeat with its existing API key
    before saving the new URL.
    """
    if not DISCOVERY_ENABLED:
        log("Auto Discovery đang tắt trong config")
        return
    while True:
        sock = None
        try:
            sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            sock.setsockopt(socket.SOL_SOCKET, socket.SO_BROADCAST, 1)
            sock.bind(("0.0.0.0", DISCOVERY_PORT))
            log(f"Auto Discovery sẵn sàng: UDP {DISCOVERY_PORT}")
            while True:
                data, addr = sock.recvfrom(2048)
                if data.strip() != DISCOVERY_MAGIC:
                    continue
                payload = json.dumps(discovery_payload(), ensure_ascii=False).encode("utf-8")
                sock.sendto(payload, addr)
        except Exception as exc:
            log(f"Auto Discovery lỗi: {exc}; thử lại sau 5 giây")
            time.sleep(5)
        finally:
            if sock is not None:
                try:
                    sock.close()
                except Exception:
                    pass


def xlsx_bytes(days=7, date_value="", start_value="", end_value="") -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    if start_value and end_value:
        orders = query_orders_range(start_value, end_value)
    else:
        orders = query_orders_date(date_value) if date_value else query_orders(days=days)
    orders = sorted(orders, key=lambda o: to_vietnam_datetime(o.get("received_at", "")))
    wb = Workbook()
    ws = wb.active
    ws.title = "DonHang"
    headers = ["STT", "Mã đơn hàng", "SĐT", "Thời gian nhận đơn"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
        cell.alignment = Alignment(horizontal="center")
    for i, o in enumerate(orders, 1):
        code = o.get("business_order_code") or business_order_code(o.get("platform", ""), o.get("order_code", ""), o.get("display_order_id", "")) or "Đang xác định"
        ws.append([i, code, normalize_phone(o.get("phone", "")), display_timestamp(o.get("received_at", ""))])
    widths = [7, 20, 18, 28]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


from app.view.loader import load_dashboard_html
DASHBOARD_HTML = load_dashboard_html()


class HubHandler(BaseHTTPRequestHandler):
    server_version = "OrderRecorderHub/2.2.5"

    def log_message(self, fmt, *args):
        log(f"{self.client_address[0]} {fmt % args}")

    def _cors(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type, X-Order-Recorder-Key")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")

    def send_json(self, data, status=200):
        raw = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(status)
        self._cors()
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def send_text(self, text, status=200, content_type="text/plain; charset=utf-8"):
        raw = text.encode("utf-8")
        self.send_response(status)
        self._cors()
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def is_loopback(self):
        return dashboard_ip_allowed(self.client_address[0]) and self.client_address[0] in {"127.0.0.1", "::1"}

    def dashboard_reader_allowed(self):
        return dashboard_ip_allowed(self.client_address[0])

    def authorized_writer(self):
        if self.is_loopback():
            return True
        return bool(API_KEY) and secrets.compare_digest(self.headers.get("X-Order-Recorder-Key", ""), API_KEY)

    def do_OPTIONS(self):
        self.send_response(204)
        self._cors()
        self.end_headers()

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        path = u.path
        qs = urllib.parse.parse_qs(u.query)
        if path == "/":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Dashboard chỉ cho phép từ PC Hub hoặc thiết bị trong Tailscale"}, 403)
            return self.send_text(DASHBOARD_HTML, content_type="text/html; charset=utf-8")
        if path == "/health" or path == "/api/health":
            return self.send_json({"ok": True, "app": APP_NAME, "version": APP_VERSION, "time": local_now_iso()})
        if path == "/api/discovery":
            # Public but intentionally contains no API key or order data.
            # Used as a TCP fallback when UDP broadcast is blocked by a firewall/AP.
            return self.send_json(discovery_payload())
        if path == "/api/stats":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Dữ liệu chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            date_value = (qs.get("date") or [""])[0]
            start_value = (qs.get("start_date") or [""])[0]
            end_value = (qs.get("end_date") or [""])[0]
            try:
                if start_value and end_value:
                    return self.send_json(stats_for_range(start_value, end_value))
                return self.send_json(stats_for_date(date_value) if date_value else stats_today())
            except ValueError as exc:
                return self.send_json({"error": str(exc)}, 400)
        if path == "/api/stats/today":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Dữ liệu chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            return self.send_json(stats_today())
        if path == "/api/orders":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Dữ liệu chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            try:
                days = int((qs.get("days") or ["1"])[0])
            except Exception:
                days = 1
            platform = platform_name((qs.get("platform") or [""])[0])
            q = (qs.get("q") or [""])[0]
            date_value = (qs.get("date") or [""])[0]
            start_value = (qs.get("start_date") or [""])[0]
            end_value = (qs.get("end_date") or [""])[0]
            try:
                if start_value and end_value:
                    orders = query_orders_range(start_value, end_value, platform, q)
                else:
                    orders = query_orders_date(date_value, platform, q) if date_value else query_orders(days, platform, q)
                return self.send_json({"orders": orders})
            except ValueError as exc:
                return self.send_json({"error": str(exc)}, 400)
        if path == "/api/reconciliation":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Đối soát chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            start_value = (qs.get("start_date") or [""])[0]
            end_value = (qs.get("end_date") or [""])[0]
            try:
                if not start_value or not end_value:
                    today=datetime.now(VIETNAM_TZ).date().isoformat();start_value=end_value=today
                return self.send_json(reconciliation_for_range(start_value,end_value))
            except ValueError as exc:
                return self.send_json({"error": str(exc)}, 400)
        if path == "/api/trash":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Thùng rác chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            try: limit=int((qs.get("limit") or ["100"])[0])
            except Exception: limit=100
            return self.send_json({"orders": get_trash(limit,(qs.get("q") or [""])[0])})
        if path == "/api/order-detail":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Chi tiết đơn chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            try:
                return self.send_json(get_order_detail((qs.get("id") or ["0"])[0]))
            except ValueError as exc:
                return self.send_json({"error": str(exc)}, 404)
        if path == "/api/device-status":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Trạng thái chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            return self.send_json(device_status())
        if path == "/api/audit-logs":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Nhật ký chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            try:
                limit = int((qs.get("limit") or ["100"])[0])
            except Exception:
                limit = 100
            return self.send_json({"logs": get_audit_logs(limit)})
        if path == "/api/local-config":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Config chỉ xem từ PC Hub hoặc qua Tailscale"}, 403)
            remote = not self.is_loopback()
            return self.send_json({
                "lan_ip": get_local_ip(),
                "port": PORT,
                "api_key": "" if remote else API_KEY,
                "api_key_hidden": remote,
                "db_path": str(DB_PATH),
                "retention_days": RETENTION_DAYS,
                "access": "tailscale" if remote else "local",
                "hub_id": HUB_ID,
                "discovery_enabled": DISCOVERY_ENABLED,
                "discovery_port": DISCOVERY_PORT,
                "admin_password_set": admin_password_is_set(),
                "database_backup": database_backup_status(),
                "pc_alerts": {"enabled": PC_ALERTS_ENABLED, "sound": PC_ALERT_SOUND, "native": os.name == "nt"},
            })
        if path == "/api/export.xlsx":
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Excel chỉ xuất từ PC Hub hoặc qua Tailscale"}, 403)
            try:
                date_value = (qs.get("date") or [""])[0]
                start_value = (qs.get("start_date") or [""])[0]
                end_value = (qs.get("end_date") or [""])[0]
                days = int((qs.get("days") or [str(RETENTION_DAYS)])[0])
                data = xlsx_bytes(days, date_value, start_value, end_value)
                if start_value and end_value:
                    start_label = datetime.strptime(start_value, '%Y-%m-%d').strftime('%d-%m-%Y')
                    end_label = datetime.strptime(end_value, '%Y-%m-%d').strftime('%d-%m-%Y')
                    filename = f"OrderRecorder_{start_label}.xlsx" if start_value == end_value else f"OrderRecorder_{start_label}_den_{end_label}.xlsx"
                else:
                    label = date_value or datetime.now(VIETNAM_TZ).strftime('%Y-%m-%d')
                    try:
                        out_label = datetime.strptime(label, '%Y-%m-%d').strftime('%d-%m-%Y')
                    except Exception:
                        out_label = label
                    filename = f"OrderRecorder_{out_label}.xlsx"
                self.send_response(200)
                self._cors()
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)
                return
            except Exception as exc:
                log(traceback.format_exc())
                return self.send_json({"error": str(exc)}, 500)
        return self.send_json({"error": "Not found"}, 404)

    def do_POST(self):
        u = urllib.parse.urlparse(self.path)
        dashboard_paths = {"/api/manual-backup", "/api/order-edit", "/api/order-delete", "/api/order-restore", "/api/database-backup", "/api/import-xlsx", "/api/alert-test"}
        writer_paths = {"/api/orders", "/api/heartbeat", "/api/order-events"}
        special_paths = {"/api/admin-password"}
        allowed = dashboard_paths | writer_paths | special_paths
        if u.path not in allowed:
            return self.send_json({"error": "Not found"}, 404)

        if u.path in dashboard_paths:
            if not self.dashboard_reader_allowed():
                return self.send_json({"error": "Thao tác quản lý chỉ cho phép từ PC Hub hoặc qua Tailscale"}, 403)
        elif u.path in writer_paths:
            if not self.authorized_writer():
                return self.send_json({"error": "API key không đúng"}, 401)
        elif u.path == "/api/admin-password" and not self.is_loopback():
            return self.send_json({"error": "Mật khẩu admin chỉ được đặt/đổi trực tiếp trên PC Hub"}, 403)

        try:
            n = int(self.headers.get("Content-Length", "0"))
            max_body = 8 * 1024 * 1024 if u.path == "/api/import-xlsx" else 128 * 1024
            if n <= 0 or n > max_body:
                raise ValueError("Request body không hợp lệ hoặc vượt giới hạn")
            payload = json.loads(self.rfile.read(n).decode("utf-8"))
            source_ip = str(self.client_address[0] or "")

            if u.path == "/api/heartbeat":
                return self.send_json({"ok": True, "device": upsert_heartbeat(payload)}, 200)
            if u.path == "/api/order-events":
                if isinstance(payload, dict) and isinstance(payload.get("events"), list):
                    saved=[upsert_order_event(x) for x in payload.get("events")[:500]]
                    return self.send_json({"ok": True, "events": saved, "count": len(saved)}, 200)
                return self.send_json({"ok": True, "event": upsert_order_event(payload)}, 200)
            if u.path == "/api/manual-backup":
                row, action = manual_backup_order(payload)
                return self.send_json({"ok": True, "action": action, "order": row}, 201 if action == "created" else 200)
            if u.path == "/api/order-edit":
                row = edit_order(payload, source_ip)
                return self.send_json({"ok": True, "order": row}, 200)
            if u.path == "/api/order-delete":
                result = soft_delete_order(payload, source_ip)
                return self.send_json({"ok": True, "result": result}, 200)
            if u.path == "/api/order-restore":
                row = restore_order(payload, source_ip)
                return self.send_json({"ok": True, "order": row}, 200)
            if u.path == "/api/database-backup":
                result=create_database_backup(force=True)
                if not result.get("ok"):
                    return self.send_json({"error": result.get("error") or "Không sao lưu được database"}, 500)
                return self.send_json(result, 200)
            if u.path == "/api/import-xlsx":
                result = import_xlsx_orders(payload, source_ip)
                return self.send_json(result, 200)
            if u.path == "/api/alert-test":
                event={"platform":"shopeefood","event_type":"risk","order_code":"SPF-TEST","short_order_number":"TEST",
                       "event_at":local_now_iso(),"reason":"Cảnh báo thử từ Hub"}
                dispatch_pc_alert(event, force=True)
                return self.send_json({"ok": True, "message": "Đã gửi cảnh báo thử tới PC Hub"}, 200)
            if u.path == "/api/admin-password":
                result = set_admin_password(payload)
                return self.send_json(result, 200)

            row, created = upsert_order(payload)
            cleanup_old_orders()
            return self.send_json({"ok": True, "created": created, "order": row}, 201 if created else 200)
        except ManualBackupConflict as exc:
            return self.send_json({"error": str(exc)}, 409)
        except PermissionError as exc:
            return self.send_json({"error": str(exc)}, 403)
        except ValueError as exc:
            return self.send_json({"error": str(exc)}, 400)
        except sqlite3.IntegrityError as exc:
            return self.send_json({"error": "Dữ liệu trùng", "detail": str(exc)}, 409)
        except Exception as exc:
            log(traceback.format_exc())
            return self.send_json({"error": str(exc)}, 500)


def periodic_cleanup():
    while True:
        time.sleep(3600)
        removed = cleanup_old_orders()
        if removed:
            log(f"Đã xóa {removed} đơn quá {RETENTION_DAYS} ngày")
        if BACKUP_ENABLED:
            create_database_backup(force=False)


def main():
    init_db()
    ip = get_local_ip()
    log(f"{APP_NAME} v{APP_VERSION}")
    log(f"Dashboard PC: http://127.0.0.1:{PORT}")
    log(f"Địa chỉ LAN: http://{ip}:{PORT}")
    log("Dashboard remote: cho phép localhost + Tailscale (100.64.0.0/10, fd7a:115c:a1e0::/48)")
    log(f"Hub ID: {HUB_ID}")
    log(f"Auto Discovery: {'BẬT' if DISCOVERY_ENABLED else 'TẮT'} · UDP {DISCOVERY_PORT}")
    log(f"Database: {DB_PATH}")
    log(f"Lưu dữ liệu: {RETENTION_DAYS} ngày")
    log("Manual Backup: BẬT · chỉ PC Hub/Tailscale được phép chèn SĐT")
    log("Reconciliation telemetry: BẬT · seen/captured/risk/miss")
    log(f"PC Early Risk Alert: {'BẬT' if PC_ALERTS_ENABLED else 'TẮT'} · native Windows + âm thanh + banner Dashboard")
    log("XLSX Import: BẬT · bắt buộc mật khẩu admin + backup database trước import")
    log(f"Database Backup: {'BẬT' if BACKUP_ENABLED else 'TẮT'} · giữ {BACKUP_RETENTION_DAYS} ngày")
    if BACKUP_ENABLED:
        create_database_backup(force=False)
    print("\nĐể SUNMI gửi dữ liệu, cần cùng Wi-Fi với PC và dùng API key hiển thị trên Dashboard.\n", flush=True)
    threading.Thread(target=periodic_cleanup, daemon=True, name="hub-cleanup").start()
    if DISCOVERY_ENABLED:
        threading.Thread(target=discovery_server, daemon=True, name="hub-discovery").start()
    httpd = ThreadingHTTPServer(("0.0.0.0", PORT), HubHandler)
    background = "--background" in sys.argv
    if CONFIG.get("auto_open_dashboard", True) and not background:
        threading.Timer(1.0, lambda: webbrowser.open(f"http://127.0.0.1:{PORT}")).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        log("Đang tắt Hub...")
    finally:
        httpd.server_close()


if __name__ == "__main__":
    main()
